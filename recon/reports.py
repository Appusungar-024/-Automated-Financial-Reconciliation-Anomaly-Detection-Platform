"""
Audit report: one Excel workbook an auditor or finance controller can file.

Contents -- Summary (run facts, headline numbers, control-total checks, exposure by type, ageing,
data quality, parameters), Exceptions (with workflow state), Judgmental Matches (the tolerance /
fuzzy / manual decisions someone may want to challenge), Anomalies, Data Quality and the full
Audit Trail.

Deliberately static values, not formulas: this is a signed-off snapshot of a database at a moment
in time, and it should not change if someone edits a cell. (Live analysis happens in the dashboard
and the SQL views.)

Security: exports contain text copied from user-supplied files. A vendor named ``=HYPERLINK(...)``
would execute when opened in Excel (CSV / formula injection), so every string cell that could be
read as a formula is neutralised with a leading apostrophe.
"""
from __future__ import annotations

import io
import json
from datetime import datetime

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import config, repository

FONT = "Arial"
HEADER_FILL = PatternFill("solid", start_color="1F2937", end_color="1F2937")
BAND_FILL = {"CRITICAL": "FECACA", "HIGH": "FED7AA", "MEDIUM": "FEF08A"}
PASS_FILL, FAIL_FILL = PatternFill("solid", start_color="BBF7D0", end_color="BBF7D0"), PatternFill("solid", start_color="FCA5A5", end_color="FCA5A5")
MONEY = f'"{config.CURRENCY_SYMBOL}"#,##0.00'
_FORMULA_STARTS = ("=", "+", "-", "@", "\t", "\r")


def safe_cell(value):
    """Neutralise anything Excel would interpret as a formula."""
    if isinstance(value, str) and value.startswith(_FORMULA_STARTS):
        return "'" + value
    return value


def _py(value):
    """pandas / numpy scalars -> plain Python values openpyxl accepts."""
    if value is None or (not isinstance(value, (str, list, dict)) and pd.isna(value)):
        return None
    if isinstance(value, pd.Timestamp):
        return value.to_pydatetime()
    if hasattr(value, "item"):
        return value.item()
    return value


def _write_table(ws, df: pd.DataFrame, spec: list, start_row: int = 1, band_col: str | None = None):
    """spec: (column, header, width, number_format|None)."""
    cols = [(c, h, w, f) for c, h, w, f in spec if c in df.columns]
    for j, (_, header, width, _fmt) in enumerate(cols, start=1):
        cell = ws.cell(row=start_row, column=j, value=header)
        cell.font, cell.fill = Font(name=FONT, bold=True, color="FFFFFF"), HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(j)].width = width
    body = Font(name=FONT)
    for i, row in enumerate(df[[c for c, *_ in cols]].itertuples(index=False), start=start_row + 1):
        for j, ((name, _h, _w, fmt), value) in enumerate(zip(cols, row), start=1):
            cell = ws.cell(row=i, column=j, value=safe_cell(_py(value)))
            cell.font = body
            if fmt:
                cell.number_format = fmt
            if band_col and name == band_col and value in BAND_FILL:
                cell.fill = PatternFill("solid", start_color=BAND_FILL[value], end_color=BAND_FILL[value])
    ws.freeze_panes = ws.cell(row=start_row + 1, column=1)
    if len(df):
        ws.auto_filter.ref = f"A{start_row}:{get_column_letter(len(cols))}{start_row + len(df)}"
    return start_row + len(df) + 1


def _title(ws, row, text, size=12):
    ws.cell(row=row, column=1, value=text).font = Font(name=FONT, bold=True, size=size)
    return row + 1


def _kv(ws, row, label, value, fmt=None):
    ws.cell(row=row, column=1, value=label).font = Font(name=FONT, bold=True)
    cell = ws.cell(row=row, column=2, value=safe_cell(_py(value)))
    cell.font, cell.alignment = Font(name=FONT), Alignment(horizontal="left")
    if fmt:
        cell.number_format = fmt
    return row + 1


EXCEPTION_SPEC = [
    ("exception_id", "ID", 8, None), ("priority_band", "Priority", 11, None), ("priority_score", "Score", 8, "0.0"),
    ("exception_type", "Type", 17, None), ("status", "Status", 12, None), ("assigned_to", "Assigned to", 14, None),
    ("source_system", "Source", 16, None), ("txn_date", "Date", 12, "yyyy-mm-dd"), ("vendor_raw", "Vendor / narration", 34, None),
    ("reference", "Reference", 18, None), ("amount", "Amount", 16, MONEY), ("age_days", "Age (days)", 10, "0"),
    ("risk_score", "Risk", 8, "0.0"), ("suggested_match_row_id", "Suggested match", 18, None),
    ("suggested_match_confidence", "Confidence", 11, "0.0"), ("suggested_match_reason", "Reason", 44, None),
    ("resolution_code", "Resolution", 20, None), ("resolved_by", "Resolved by", 14, None),
    ("resolved_at", "Resolved at", 18, "yyyy-mm-dd hh:mm"), ("resolution_notes", "Notes", 40, None), ("row_id", "Row id", 18, None)]

MATCH_SPEC = [
    ("match_id", "ID", 8, None), ("pair_name", "Pair", 18, None), ("match_type", "Rule", 18, None),
    ("match_confidence", "Confidence", 11, "0.0"), ("date_a", "Date A", 12, "yyyy-mm-dd"), ("date_b", "Date B", 12, "yyyy-mm-dd"),
    ("date_diff_days", "Date diff (d)", 12, "0"), ("amount_a", "Amount A", 16, MONEY), ("amount_b", "Amount B", 16, MONEY),
    ("amount_diff", "Amount diff", 14, MONEY), ("vendor_similarity", "Vendor sim.", 11, "0.0"), ("vendor_a", "Vendor A", 30, None),
    ("vendor_b", "Vendor B", 30, None), ("rule_applied", "Why matched", 52, None), ("decided_by", "Decided by", 14, None),
    ("matched_at", "Matched at", 18, "yyyy-mm-dd hh:mm")]

ANOMALY_SPEC = [
    ("row_id", "Row id", 18, None), ("source_system", "Source", 16, None), ("txn_date", "Date", 12, "yyyy-mm-dd"),
    ("vendor_raw", "Vendor / narration", 34, None), ("amount", "Amount", 16, MONEY), ("risk_score", "Risk", 8, "0.0"),
    ("rule_score", "Rules", 8, "0"), ("ml_score", "Model", 8, "0.0"), ("suspicion_reasons", "Why flagged", 90, None)]

FAILURE_SPEC = [
    ("source_system", "Source", 16, None), ("file_name", "File", 24, None), ("source_row", "Row", 8, "0"),
    ("severity", "Severity", 10, None), ("code", "Code", 22, None), ("field", "Field", 14, None),
    ("raw_value", "Value", 24, None), ("message", "Message", 60, None), ("row_excluded", "Excluded", 9, None)]

AUDIT_SPEC = [
    ("audit_id", "ID", 8, None), ("event_time", "Time (UTC)", 20, "yyyy-mm-dd hh:mm:ss"), ("event_type", "Event", 26, None),
    ("actor", "Actor", 14, None), ("entity_type", "Entity", 11, None), ("entity_id", "Entity id", 16, None),
    ("old_value", "From", 14, None), ("new_value", "To", 14, None), ("details", "Details", 90, None)]


def exceptions_workbook(df: pd.DataFrame, title: str = "Exceptions") -> bytes:
    """A formatted export of any (e.g. filtered) set of exceptions."""
    wb = Workbook()
    ws = wb.active
    ws.title = title
    _write_table(ws, df, EXCEPTION_SPEC, band_col="priority_band")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_audit_report(run_id: str, actor: str = "system") -> bytes:
    k = repository.kpis(run_id)
    if not k:
        raise ValueError(f"unknown run {run_id!r}")
    run, sym = k["run"], config.CURRENCY_SYMBOL
    wb = Workbook()

    # ------------------------------------------------------------------ Summary
    ws = wb.active
    ws.title = "Summary"
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 52, 26
    for col in "CDEF":
        ws.column_dimensions[col].width = 18
    r = 1
    ws.cell(row=r, column=1, value="Reconciliation audit report").font = Font(name=FONT, bold=True, size=16)
    r += 2
    r = _title(ws, r, "Run")
    for label, value, fmt in [("Run id", run["run_id"], None), ("Status", run["status"], None), ("Notes", run["notes"] or "", None),
                              ("Started (UTC)", run["started_at"], "yyyy-mm-dd hh:mm"), ("Finished (UTC)", run["finished_at"], "yyyy-mm-dd hh:mm"),
                              ("As-of date (ageing)", run["as_of_date"], "yyyy-mm-dd"), ("Report generated (UTC)", datetime.utcnow(), "yyyy-mm-dd hh:mm"),
                              ("Report generated by", actor, None)]:
        r = _kv(ws, r, label, value, fmt)
    r += 1
    r = _title(ws, r, "Headline numbers (live: reflects reviewer actions since the run)")
    for label, value, fmt in [
            ("Transactions reconciled", k["total_rows"], "#,##0"), ("Matched (rows)", k["matched_rows"], "#,##0"),
            ("Reconciliation rate by count (%)", round(k["match_rate"], 2), "0.00"),
            ("Reconciliation rate by value (%)", round(k["value_match_rate"], 2), "0.00"),
            ("Exceptions raised", k["exceptions_total"], "#,##0"), ("Exceptions still open", k["exceptions_open"], "#,##0"),
            ("Resolution rate (%)", round(k["resolution_rate"], 2), "0.00"),
            ("Unresolved value, gross (a mismatched pair counts both sides)", k["unresolved_amount"], MONEY),
            ("Average age of open exceptions (days)", round(k["avg_age_open"], 1), "0.0"),
            ("Open CRITICAL / HIGH items", k["urgent_open"], "#,##0"),
            ("Transactions flagged suspicious", k["suspicious"], "#,##0"),
            ("Source rows quarantined at ingestion", k["quarantined_rows"], "#,##0"),
            ("Audit events recorded", k["audit_events"], "#,##0")]:
        r = _kv(ws, r, label, value, fmt)
    r += 1
    r = _title(ws, r, "Control totals")
    for j, h in enumerate(["Check", "Expected", "Actual", "Result"], start=1):
        c = ws.cell(row=r, column=j, value=h)
        c.font, c.fill = Font(name=FONT, bold=True, color="FFFFFF"), HEADER_FILL
    r += 1
    for chk in repository.control_checks(run_id):
        ws.cell(row=r, column=1, value=chk["check"]).font = Font(name=FONT)
        ws.cell(row=r, column=2, value=chk["expected"]).number_format = "#,##0"
        ws.cell(row=r, column=3, value=chk["actual"]).number_format = "#,##0"
        res = ws.cell(row=r, column=4, value="PASS" if chk["passed"] else "FAIL")
        res.font, res.fill = Font(name=FONT, bold=True), PASS_FILL if chk["passed"] else FAIL_FILL
        r += 1
    r += 1
    summary = repository.exception_summary(run_id)
    r = _title(ws, r, "Exposure by exception type")
    if not summary.empty:
        by_type = summary.groupby("exception_type")[["exception_count", "unresolved_count", "total_amount", "unresolved_amount"]].sum().reset_index()
        r = _write_table(ws, by_type, [("exception_type", "Type", 52, None), ("exception_count", "Exceptions", 26, "#,##0"),
                                       ("unresolved_count", "Unresolved", 18, "#,##0"), ("total_amount", "Total value", 18, MONEY),
                                       ("unresolved_amount", "Unresolved value", 18, MONEY)], start_row=r)
        ws.freeze_panes = None
        ws.auto_filter.ref = None
    r += 1
    aging = repository.exception_aging(run_id)
    r = _title(ws, r, "Ageing of unresolved exceptions")
    if not aging.empty:
        r = _write_table(ws, aging, [("aging_bucket", "Bucket", 52, None), ("exception_count", "Exceptions", 26, "#,##0"),
                                     ("unresolved_amount", "Value", 18, MONEY)], start_row=r)
        ws.freeze_panes = None
        ws.auto_filter.ref = None
    r += 1
    r = _title(ws, r, "Source files")
    files = repository.source_files_df(run_id)
    if not files.empty:
        r = _write_table(ws, files, [("source_system", "Source", 52, None), ("file_name", "File", 26, None), ("rows_read", "Rows read", 18, "#,##0"),
                                     ("rows_accepted", "Accepted", 18, "#,##0"), ("rows_excluded", "Excluded", 18, "#,##0"),
                                     ("sha256", "SHA-256", 66, None)], start_row=r)
        ws.freeze_panes = None
        ws.auto_filter.ref = None
    r += 1
    r = _title(ws, r, "Parameters in force for this run")
    snapshot = json.loads(run["config_snapshot"] or "{}")
    for key in ("AMOUNT_TOLERANCE_ABS", "AMOUNT_TOLERANCE_PCT", "FUZZY_DATE_WINDOW_DAYS", "FUZZY_VENDOR_THRESHOLD",
                "ZSCORE_OUTLIER_THRESHOLD", "RISK_FLAG_THRESHOLD", "ML_N_ESTIMATORS", "ML_RANDOM_STATE"):
        if key in snapshot:
            r = _kv(ws, r, key, str(snapshot[key]))
    r += 2
    for label in ("Prepared by:", "Reviewed by:", "Date:"):
        ws.cell(row=r, column=1, value=label).font = Font(name=FONT, bold=True)
        ws.cell(row=r, column=2, value="______________________________")
        r += 2

    # ------------------------------------------------------------------ detail sheets
    ws = wb.create_sheet("Exceptions")
    _write_table(ws, repository.review_queue(run_id, limit=None), EXCEPTION_SPEC, band_col="priority_band")
    ws = wb.create_sheet("Judgmental Matches")
    _write_table(ws, repository.matches(run_id, judgmental_only=True, limit=None), MATCH_SPEC)
    ws = wb.create_sheet("Anomalies")
    _write_table(ws, repository.top_risky(run_id, flagged_only=True, limit=2000), ANOMALY_SPEC)
    ws = wb.create_sheet("Data Quality")
    _write_table(ws, repository.validation_failures_df(run_id, limit=20000), FAILURE_SPEC)
    ws = wb.create_sheet("Audit Trail")
    trail = repository.audit_trail(run_id, limit=50000).sort_values("audit_id")
    _write_table(ws, trail, AUDIT_SPEC)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
