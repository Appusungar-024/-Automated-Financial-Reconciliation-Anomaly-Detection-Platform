import io

import pytest
from openpyxl import load_workbook

from recon import pipeline, reports, repository
from tests.conftest import BANK_HEADER, LEDGER_HEADER, VALIDATION_AS_OF, write_rows


def sheet_values(wb, name):
    return [[c.value for c in row] for row in wb[name].iter_rows()]


# ---------------------------------------------------------------- control totals

def test_control_totals_pass_on_a_fresh_run(mini_run):
    checks = repository.control_checks(mini_run.run_id)
    assert len(checks) == 5 and all(c["passed"] for c in checks), checks


def test_control_totals_still_balance_after_reviewers_act(mini_run):
    run_id = mini_run.run_id
    a, b = (repository.exception_for_row(run_id, r) for r in ("BANK-0000002", "LEDGER-0000003"))
    repository.manual_match(a["exception_id"], b["exception_id"], actor="asha")          # adds match rows, keeps exceptions
    c = repository.exception_for_row(run_id, "BANK-0000001")
    repository.update_exception(c["exception_id"], actor="asha", status="DISMISSED", resolution_code="FALSE_POSITIVE", note="n")
    assert all(x["passed"] for x in repository.control_checks(run_id))


def test_a_corrupted_database_fails_the_controls(mini_run):
    from sqlalchemy import text
    from recon import db
    with db.get_engine().begin() as conn:       # simulate a tampered table: an open exception whose row is also matched
        conn.execute(text("UPDATE exceptions SET row_id = 'BANK-0000000' WHERE row_id = 'BANK-0000002'"))
    failed = [c["check"] for c in repository.control_checks(mini_run.run_id) if not c["passed"]]
    assert any("both matched and still an open exception" in f for f in failed)


# ---------------------------------------------------------------- the workbook

def test_audit_report_structure_and_content(mini_run):
    run_id = mini_run.run_id
    e = repository.exception_for_row(run_id, "BANK-0000002")
    repository.update_exception(e["exception_id"], actor="asha", status="RESOLVED", resolution_code="BANK_ERROR", note="called bank")
    wb = load_workbook(io.BytesIO(reports.build_audit_report(run_id, actor="controller")))
    assert wb.sheetnames == ["Summary", "Exceptions", "Judgmental Matches", "Anomalies", "Data Quality", "Audit Trail"]

    summary = [v for row in sheet_values(wb, "Summary") for v in row if v is not None]
    assert run_id in summary and "controller" in summary and "PASS" in summary and "FAIL" not in summary
    exceptions = sheet_values(wb, "Exceptions")
    assert len(exceptions) == 1 + 4                                     # header + four exceptions
    header = exceptions[0]
    row = next(r for r in exceptions[1:] if r[header.index("Row id")] == "BANK-0000002")
    assert row[header.index("Status")] == "RESOLVED" and row[header.index("Resolved by")] == "asha"
    trail = sheet_values(wb, "Audit Trail")
    assert any("asha" in r for r in trail) and any("RUN_COMPLETED" in r for r in trail)


def test_report_amounts_are_numbers_with_currency_format_not_text(mini_run):
    wb = load_workbook(io.BytesIO(reports.build_audit_report(mini_run.run_id)))
    ws = wb["Exceptions"]
    header = [c.value for c in ws[1]]
    cell = ws.cell(row=2, column=header.index("Amount") + 1)
    assert isinstance(cell.value, (int, float)) and "#,##0.00" in cell.number_format


def test_priority_bands_are_colour_coded(mini_run):
    ws = load_workbook(io.BytesIO(reports.build_audit_report(mini_run.run_id)))["Exceptions"]
    header = [c.value for c in ws[1]]
    fills = {ws.cell(row=r, column=header.index("Priority") + 1).value:
             ws.cell(row=r, column=header.index("Priority") + 1).fill.start_color.rgb for r in range(2, ws.max_row + 1)}
    assert any(band in fills for band in ("CRITICAL", "HIGH", "MEDIUM", "LOW"))


def test_unknown_run_is_rejected(tmp_db):
    with pytest.raises(ValueError, match="unknown run"):
        reports.build_audit_report("nope")


# ---------------------------------------------------------------- formula injection

@pytest.mark.parametrize("text", ["=HYPERLINK(\"http://evil\",\"x\")", "+SUM(A1)", "-2+3", "@cmd", "\t=1"])
def test_formula_like_strings_are_neutralised(text):
    assert reports.safe_cell(text) == "'" + text


@pytest.mark.parametrize("value", ["Acme Ltd", "", 42, 3.5, None])
def test_ordinary_values_pass_through_untouched(value):
    assert reports.safe_cell(value) == value


def test_malicious_vendor_names_cannot_become_live_formulas_in_an_export(tmp_db, tmp_path):
    bank = write_rows(tmp_path / "bank.csv", [BANK_HEADER, ["12/08/2026", "", "NEFT-OK-1", "B1", "100.00", ""]])
    ledger = write_rows(tmp_path / "ledger.csv", [LEDGER_HEADER,
        ["2026-08-12", "OK", "V1", "100.00", "", ""],
        ["2026-08-13", "=cmd|' /C calc'!A0", "V2", "55.00", "", ""]])
    run = pipeline.run({"BANK": bank, "LEDGER": ledger}, validation_as_of=VALIDATION_AS_OF)
    wb = load_workbook(io.BytesIO(reports.exceptions_workbook(repository.review_queue(run.run_id))))
    ws = wb.active
    header = [c.value for c in ws[1]]
    vendor = ws.cell(row=2, column=header.index("Vendor / narration") + 1)
    assert vendor.value.startswith("'=cmd") and vendor.data_type == "s"
